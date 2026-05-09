import io
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

class FileProcessor:
    @staticmethod
    def extract_text(file_content: bytes, filename: str) -> str:
        extension = filename.split('.')[-1].lower()

        if extension == 'pdf':
            return FileProcessor._parse_pdf(file_content)
        elif extension in ['docx', 'doc']:
            return FileProcessor._parse_docx(file_content)
        elif extension in ['xlsx', 'xls']:
            return FileProcessor._parse_xlsx(file_content)
        elif extension == 'txt':
            return file_content.decode('utf-8')
        else:
            raise ValueError(f"Unsupported file extension: {extension}")

    @staticmethod
    def _parse_pdf(content: bytes) -> str:
        reader = PdfReader(io.BytesIO(content))
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text

    @staticmethod
    def _parse_docx(content: bytes) -> str:
        doc = Document(io.BytesIO(content))
        return "\n".join([para.text for para in doc.paragraphs])

    @staticmethod
    def _parse_xlsx(content: bytes) -> str:
        wb = load_workbook(io.BytesIO(content), data_only=True)
        text_parts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                # Объединяем ячейки строки в текст
                row_text = " ".join([str(cell) for cell in row if cell is not None])
                if row_text.strip():
                    text_parts.append(row_text)
        return "\n".join(text_parts)

file_processor = FileProcessor()
